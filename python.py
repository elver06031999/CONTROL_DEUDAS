import os
import re
import shutil
import sqlite3
import subprocess
import tkinter as tk
from datetime import datetime, timedelta
from tkinter import filedialog, messagebox, ttk

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    OPENPYXL_DISPONIBLE = True
except ImportError:
    OPENPYXL_DISPONIBLE = False

# Carpetas del sistema
CARPETA_VOUCHERS = "vouchers"
if not os.path.exists(CARPETA_VOUCHERS):
    os.makedirs(CARPETA_VOUCHERS)


# =====================================================================
# 1. BASE DE DATOS Y UTILIDADES DE VALIDACIÓN
# =====================================================================
def conectar_bd():
    conn = sqlite3.connect("control_deudas.db")
    conn.execute("PRAGMA foreign_keys = ON;")  # Habilitar Claves Foráneas
    return conn


def inicializar_bd():
    conn = conectar_bd()
    cursor = conn.cursor()
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS clientes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        ruc TEXT UNIQUE NOT NULL,
        razon_social TEXT NOT NULL,
        contacto TEXT,
        regimen TEXT DEFAULT 'MYPE Tributario'
    )
    """)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS obligaciones (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        cliente_id INTEGER,
        tipo TEXT NOT NULL,
        periodo TEXT NOT NULL,
        monto REAL NOT NULL,
        fecha_vencimiento TEXT,
        estado TEXT DEFAULT 'PENDIENTE',
        fecha_pago TEXT,
        situacion TEXT DEFAULT 'Declarado',
        ruta_voucher TEXT,
        descripcion_tributos TEXT,
        FOREIGN KEY (cliente_id) REFERENCES clientes(id) ON DELETE CASCADE
    )
    """)

    for sentencia in [
        "ALTER TABLE clientes ADD COLUMN regimen TEXT DEFAULT 'MYPE Tributario'",
        "ALTER TABLE obligaciones ADD COLUMN situacion TEXT DEFAULT 'Declarado'",
        "ALTER TABLE obligaciones ADD COLUMN ruta_voucher TEXT",
        "ALTER TABLE obligaciones ADD COLUMN descripcion_tributos TEXT",
    ]:
        try:
            cursor.execute(sentencia)
        except sqlite3.OperationalError:
            pass

    conn.commit()
    conn.close()


def validar_fecha_iso(fecha_str):
    """Valida formato AAAA-MM-DD"""
    if not fecha_str:
        return True
    try:
        datetime.strptime(fecha_str, "%Y-%m-%d")
        return True
    except ValueError:
        return False


def validar_periodo(periodo_str):
    """Valida formato AAAA-MM"""
    if not periodo_str:
        return False
    return bool(re.match(r"^\d{4}-(0[1-9]|1[0-2])$", periodo_str))


def calcular_vencimiento_sunat_aproximado(ruc, periodo):
    """Calcula fecha aproximada de vencimiento SUNAT según el último dígito del RUC."""
    if not ruc or len(ruc) != 11 or not validar_periodo(periodo):
        return ""
    try:
        ultimo_digito = int(ruc[-1])
        año, mes = map(int, periodo.split("-"))
        # El vencimiento de un periodo es en el mes siguiente
        if mes == 12:
            mes_venc = 1
            año_venc = año + 1
        else:
            mes_venc = mes + 1
            año_venc = año

        # Offset estándar aproximado para cronograma SUNAT según dígito
        dias_offset = {
            0: 14,
            1: 15,
            2: 16,
            3: 17,
            4: 18,
            5: 19,
            6: 20,
            7: 21,
            8: 22,
            9: 23,
        }
        dia = dias_offset.get(ultimo_digito, 15)
        return f"{año_venc:04d}-{mes_venc:02d}-{dia:02d}"
    except Exception:
        return ""


# =====================================================================
# 2. INTERFAZ GRÁFICA MODERNA
# =====================================================================
class AppControlDeudas:

    def __init__(self, root):
        self.root = root
        self.root.title("Control de Deudas y Tributos - SUNAT / AFP")
        self.root.geometry("1300x820")
        self.root.config(bg="#f8f9fa")

        inicializar_bd()
        self.configurar_estilos()
        self.crear_componentes()
        self.cargar_datos()

    def configurar_estilos(self):
        self.style = ttk.Style()
        self.style.theme_use("clam")

        self.style.configure(
            ".", font=("Segoe UI", 10), background="#f8f9fa", foreground="#2d3748"
        )
        self.style.configure("TNotebook", background="#f8f9fa", borderwidth=0)
        self.style.configure(
            "TNotebook.Tab",
            font=("Segoe UI", 10, "bold"),
            padding=[15, 6],
            background="#e2e8f0",
            foreground="#4a5568",
        )
        self.style.map(
            "TNotebook.Tab",
            background=[("selected", "#ffffff")],
            foreground=[("selected", "#2b6cb0")],
        )

        self.style.configure("TCombobox", padding=5, background="#ffffff")

        self.style.configure(
            "Treeview",
            font=("Segoe UI", 10),
            rowheight=28,
            background="#ffffff",
            fieldbackground="#ffffff",
            borderwidth=0,
        )
        self.style.configure(
            "Treeview.Heading",
            font=("Segoe UI", 10, "bold"),
            background="#edf2f7",
            foreground="#2d3748",
            relief="flat",
            padding=6,
        )
        self.style.map(
            "Treeview",
            background=[("selected", "#ebf8ff")],
            foreground=[("selected", "#2b6cb0")],
        )

    def crear_componentes(self):
        lbl_titulo = tk.Label(
            self.root,
            text="💼  CONTROL DE DEUDAS Y VOUCHERS - SUNAT / AFP",
            font=("Segoe UI", 15, "bold"),
            bg="#1a365d",
            fg="white",
            pady=12,
        )
        lbl_titulo.pack(fill=tk.X)

        frame_cuerpo = tk.Frame(self.root, bg="#f8f9fa")
        frame_cuerpo.pack(fill=tk.BOTH, expand=True, padx=15, pady=15)

        # ================= PANEL IZQUIERDO CON SCROLLBAR =================
        frame_contenedor_izq = tk.Frame(
            frame_cuerpo, bg="#ffffff", width=360, bd=1, relief="solid"
        )
        frame_contenedor_izq.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 15))
        frame_contenedor_izq.pack_propagate(False)

        canvas_izq = tk.Canvas(
            frame_contenedor_izq, bg="#ffffff", bd=0, highlightthickness=0
        )
        scrollbar_izq = ttk.Scrollbar(
            frame_contenedor_izq, orient="vertical", command=canvas_izq.yview
        )

        pad_frame = tk.Frame(canvas_izq, bg="#ffffff")
        pad_frame.bind(
            "<Configure>",
            lambda e: canvas_izq.configure(scrollregion=canvas_izq.bbox("all")),
        )

        canvas_izq.create_window((0, 0), window=pad_frame, anchor="nw", width=340)
        canvas_izq.configure(yscrollcommand=scrollbar_izq.set)

        scrollbar_izq.pack(side=tk.RIGHT, fill=tk.Y)
        canvas_izq.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        def _al_usar_rueda(event):
            canvas_izq.yview_scroll(int(-1 * (event.delta / 120)), "units")

        pad_frame.bind_enter = lambda e: canvas_izq.bind_all(
            "<MouseWheel>", _al_usar_rueda
        )
        pad_frame.bind_leave = lambda e: canvas_izq.unbind_all("<MouseWheel>")
        pad_frame.bind("<Enter>", pad_frame.bind_enter)
        pad_frame.bind("<Leave>", pad_frame.bind_leave)

        def crear_label(parent, texto, bold=False):
            font = ("Segoe UI", 9, "bold" if bold else "normal")
            color = "#1a202c" if bold else "#4a5568"
            return tk.Label(parent, text=texto, bg="#ffffff", fg=color, font=font)

        # Sección Cliente
        crear_label(pad_frame, "REGISTRAR CLIENTE", bold=True).pack(
            anchor="w", pady=(10, 8), padx=10
        )

        crear_label(pad_frame, "RUC (11 dígitos):").pack(anchor="w", padx=10)
        self.ent_ruc = tk.Entry(pad_frame, font=("Segoe UI", 10), bd=1, relief="solid")
        self.ent_ruc.pack(fill=tk.X, pady=(2, 8), padx=10)

        crear_label(pad_frame, "Razón Social:").pack(anchor="w", padx=10)
        self.ent_razon = tk.Entry(
            pad_frame, font=("Segoe UI", 10), bd=1, relief="solid"
        )
        self.ent_razon.pack(fill=tk.X, pady=(2, 8), padx=10)

        crear_label(pad_frame, "Régimen Tributario:").pack(anchor="w", padx=10)
        self.cmb_regimen = ttk.Combobox(
            pad_frame,
            values=[
                "MYPE Tributario",
                "Régimen General",
                "RER (Especial)",
                "NRUS",
            ],
            state="readonly",
        )
        self.cmb_regimen.set("MYPE Tributario")
        self.cmb_regimen.pack(fill=tk.X, pady=(2, 8), padx=10)

        crear_label(pad_frame, "Contacto:").pack(anchor="w", padx=10)
        self.ent_contacto = tk.Entry(
            pad_frame, font=("Segoe UI", 10), bd=1, relief="solid"
        )
        self.ent_contacto.pack(fill=tk.X, pady=(2, 10), padx=10)

        frame_btn_cl = tk.Frame(pad_frame, bg="#ffffff")
        frame_btn_cl.pack(fill=tk.X, pady=(0, 10), padx=10)

        btn_reg_cliente = tk.Button(
            frame_btn_cl,
            text="➕ Añadir",
            bg="#2f855a",
            fg="white",
            font=("Segoe UI", 9, "bold"),
            relief="flat",
            activebackground="#22543d",
            activeforeground="white",
            command=self.agregar_cliente,
            cursor="hand2",
            height=1,
        )
        btn_reg_cliente.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

        btn_edit_cliente = tk.Button(
            frame_btn_cl,
            text="✏️ Editar",
            bg="#dd6b20",
            fg="white",
            font=("Segoe UI", 9, "bold"),
            relief="flat",
            activebackground="#9c4221",
            activeforeground="white",
            command=self.abrir_editar_cliente,
            cursor="hand2",
            height=1,
        )
        btn_edit_cliente.pack(side=tk.RIGHT, fill=tk.X, expand=True)

        canvas_sep = tk.Canvas(
            pad_frame, height=1, bg="#e2e8f0", bd=0, highlightthickness=0
        )
        canvas_sep.pack(fill=tk.X, pady=12, padx=10)

        # Sección Deuda
        crear_label(pad_frame, "REGISTRAR OBLIGACIÓN / DEUDA", bold=True).pack(
            anchor="w", pady=(0, 8), padx=10
        )

        crear_label(pad_frame, "Seleccionar Cliente:").pack(anchor="w", padx=10)
        self.cmb_clientes = ttk.Combobox(pad_frame, state="readonly")
        self.cmb_clientes.pack(fill=tk.X, pady=(2, 8), padx=10)

        grid_deuda = tk.Frame(pad_frame, bg="#ffffff")
        grid_deuda.pack(fill=tk.X, padx=10)

        crear_label(grid_deuda, "Tipo:").grid(
            row=0, column=0, sticky="w", padx=(0, 4)
        )
        crear_label(grid_deuda, "Situación:").grid(row=0, column=1, sticky="w")

        self.cmb_tipo = ttk.Combobox(
            grid_deuda, values=["SUNAT", "AFP"], state="readonly", width=12
        )
        self.cmb_tipo.set("SUNAT")
        self.cmb_tipo.grid(row=1, column=0, sticky="ew", pady=(2, 8), padx=(0, 6))

        self.cmb_situacion = ttk.Combobox(
            grid_deuda,
            values=["Declarado", "Por Declarar", "Fraccionado"],
            state="readonly",
            width=15,
        )
        self.cmb_situacion.set("Declarado")
        self.cmb_situacion.grid(row=1, column=1, sticky="ew", pady=(2, 8))

        crear_label(grid_deuda, "Periodo (AAAA-MM):").grid(
            row=2, column=0, sticky="w", padx=(0, 4)
        )
        crear_label(grid_deuda, "Monto (S/):").grid(row=2, column=1, sticky="w")

        self.ent_periodo = tk.Entry(
            grid_deuda, font=("Segoe UI", 10), bd=1, relief="solid", width=12
        )
        self.ent_periodo.grid(
            row=3, column=0, sticky="ew", pady=(2, 8), padx=(0, 6)
        )
        self.ent_periodo.bind("<FocusOut>", self.autocalcular_vencimiento)

        self.ent_monto = tk.Entry(
            grid_deuda, font=("Segoe UI", 10), bd=1, relief="solid", width=15
        )
        self.ent_monto.grid(row=3, column=1, sticky="ew", pady=(2, 8))

        grid_deuda.columnconfigure(0, weight=1)
        grid_deuda.columnconfigure(1, weight=1)

        crear_label(pad_frame, "Vence (AAAA-MM-DD):").pack(anchor="w", padx=10)
        self.ent_vence = tk.Entry(
            pad_frame, font=("Segoe UI", 10), bd=1, relief="solid"
        )
        self.ent_vence.pack(fill=tk.X, pady=(2, 8), padx=10)

        crear_label(pad_frame, "Fecha Pago (Opcional - AAAA-MM-DD):").pack(
            anchor="w", padx=10
        )
        self.ent_fecha_pago_reg = tk.Entry(
            pad_frame, font=("Segoe UI", 10), bd=1, relief="solid"
        )
        self.ent_fecha_pago_reg.pack(fill=tk.X, pady=(2, 8), padx=10)

        crear_label(pad_frame, "Especificación de Tributos / Detalle:").pack(
            anchor="w", padx=10
        )
        self.ent_tributos = tk.Entry(
            pad_frame, font=("Segoe UI", 10), bd=1, relief="solid"
        )
        self.ent_tributos.pack(fill=tk.X, pady=(2, 12), padx=10)

        btn_reg_deuda = tk.Button(
            pad_frame,
            text="💾  Guardar Obligación",
            bg="#2b6cb0",
            fg="white",
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            activebackground="#2c5282",
            activeforeground="white",
            command=self.agregar_obligacion,
            pady=5,
            cursor="hand2",
        )
        btn_reg_deuda.pack(fill=tk.X, padx=10, pady=(0, 15))

        # ================= PANEL DERECHO =================
        frame_derecho = tk.Frame(frame_cuerpo, bg="#f8f9fa")
        frame_derecho.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)

        frame_control = tk.Frame(
            frame_derecho,
            bg="#ffffff",
            bd=1,
            relief="solid",
            highlightthickness=0,
            highlightbackground="#e2e8f0",
        )
        frame_control.pack(fill=tk.X, pady=(0, 15))

        pad_control = tk.Frame(frame_control, bg="#ffffff", padx=12, pady=12)
        pad_control.pack(fill=tk.X)

        # Fila 1: Búsqueda global
        fila1 = tk.Frame(pad_control, bg="#ffffff")
        fila1.pack(fill=tk.X, pady=(0, 8))

        tk.Label(
            fila1,
            text="🔍  Buscar Cliente (Razón Social / RUC):",
            font=("Segoe UI", 9, "bold"),
            bg="#ffffff",
            fg="#2d3748",
        ).pack(side=tk.LEFT, padx=(0, 5))
        self.ent_buscar = tk.Entry(fila1, font=("Segoe UI", 10), bd=1, relief="solid")
        self.ent_buscar.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.ent_buscar.bind("<KeyRelease>", self.al_buscar)

        # Fila 2: Filtros específicos
        fila2 = tk.Frame(pad_control, bg="#ffffff")
        fila2.pack(fill=tk.X)

        tk.Label(
            fila2,
            text="Filtrar Empresa:",
            font=("Segoe UI", 9),
            bg="#ffffff",
            fg="#4a5568",
        ).pack(side=tk.LEFT, padx=(0, 5))
        self.cmb_filtro_ruc = ttk.Combobox(fila2, state="readonly")
        self.cmb_filtro_ruc.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 15))

        tk.Label(
            fila2,
            text="Año/Periodo:",
            font=("Segoe UI", 9),
            bg="#ffffff",
            fg="#4a5568",
        ).pack(side=tk.LEFT, padx=(0, 5))
        self.ent_filtro_periodo = tk.Entry(
            fila2, font=("Segoe UI", 10), bd=1, relief="solid", width=12
        )
        self.ent_filtro_periodo.pack(side=tk.LEFT, padx=(0, 15))

        btn_limpiar = tk.Button(
            fila2,
            text="Limpiar",
            bg="#718096",
            fg="white",
            font=("Segoe UI", 9, "bold"),
            relief="flat",
            command=self.limpiar_filtros,
            padx=15,
        )
        btn_limpiar.pack(side=tk.RIGHT)

        btn_filtro = tk.Button(
            fila2,
            text="Filtrar",
            bg="#4a5568",
            fg="white",
            font=("Segoe UI", 9, "bold"),
            relief="flat",
            command=self.cargar_datos,
            padx=15,
        )
        btn_filtro.pack(side=tk.RIGHT, padx=(0, 5))

        self.notebook = ttk.Notebook(frame_derecho)
        self.notebook.pack(fill=tk.BOTH, expand=True)

        # PESTAÑA PENDIENTES
        self.tab_pendientes = tk.Frame(self.notebook, bg="#ffffff")
        self.notebook.add(self.tab_pendientes, text="  Deudas Pendientes  ")

        frame_tabla_pend = tk.Frame(self.tab_pendientes, bg="#ffffff")
        frame_tabla_pend.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        columnas_pend = (
            "ID",
            "Razon Social",
            "RUC",
            "Régimen",
            "Tipo",
            "Situación",
            "Periodo",
            "Monto",
            "Vencimiento",
            "Tributos",
        )
        self.tabla_pendientes = ttk.Treeview(
            frame_tabla_pend, columns=columnas_pend, show="headings"
        )
        scroll_h_pend = ttk.Scrollbar(
            frame_tabla_pend,
            orient=tk.HORIZONTAL,
            command=self.tabla_pendientes.xview,
        )
        scroll_v_pend = ttk.Scrollbar(
            frame_tabla_pend,
            orient=tk.VERTICAL,
            command=self.tabla_pendientes.yview,
        )
        self.tabla_pendientes.configure(
            xscrollcommand=scroll_h_pend.set, yscrollcommand=scroll_v_pend.set
        )

        scroll_v_pend.pack(side=tk.RIGHT, fill=tk.Y)
        self.tabla_pendientes.pack(fill=tk.BOTH, expand=True)
        scroll_h_pend.pack(fill=tk.X)

        self._configurar_cabeceras_tablas(self.tabla_pendientes, columnas_pend)

        # Configuración de colores/alertas visuales para pendientes
        self.tabla_pendientes.tag_configure("vencido", background="#fed7d7")  # Rojo pastel
        self.tabla_pendientes.tag_configure("por_vencer", background="#feebc8")  # Amarillo pastel

        frame_btns_pend = tk.Frame(self.tab_pendientes, bg="#ffffff", pady=8, padx=5)
        frame_btns_pend.pack(fill=tk.X)

        tk.Button(
            frame_btns_pend,
            text="✓  Marcar como Pagado",
            bg="#2f855a",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            command=self.marcar_como_pagado,
            padx=6,
            pady=4,
        ).pack(side=tk.LEFT, padx=(0, 4))
        tk.Button(
            frame_btns_pend,
            text="✏️  Editar Deuda",
            bg="#dd6b20",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            command=self.abrir_editar_deuda_pendiente,
            padx=6,
            pady=4,
        ).pack(side=tk.LEFT, padx=4)
        tk.Button(
            frame_btns_pend,
            text="❌  Eliminar",
            bg="#e53e3e",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            command=lambda: self.eliminar_obligacion(self.tabla_pendientes),
            padx=6,
            pady=4,
        ).pack(side=tk.LEFT, padx=4)
        tk.Button(
            frame_btns_pend,
            text="📊  Exportar a Excel",
            bg="#234e52",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            command=self.exportar_pendientes,
            padx=8,
            pady=4,
        ).pack(side=tk.RIGHT)

        # PESTAÑA HISTORIAL PAGADOS
        self.tab_pagados = tk.Frame(self.notebook, bg="#ffffff")
        self.notebook.add(self.tab_pagados, text="  Historial de Pagados  ")

        frame_tabla_pag = tk.Frame(self.tab_pagados, bg="#ffffff")
        frame_tabla_pag.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        columnas_pag = (
            "ID",
            "Razon Social",
            "RUC",
            "Régimen",
            "Tipo",
            "Situación",
            "Periodo",
            "Monto",
            "Vencimiento",
            "Fecha Pago",
            "Voucher",
            "Tributos",
        )
        self.tabla_pagados = ttk.Treeview(
            frame_tabla_pag, columns=columnas_pag, show="headings"
        )
        scroll_h_pag = ttk.Scrollbar(
            frame_tabla_pag,
            orient=tk.HORIZONTAL,
            command=self.tabla_pagados.xview,
        )
        scroll_v_pag = ttk.Scrollbar(
            frame_tabla_pag, orient=tk.VERTICAL, command=self.tabla_pagados.yview
        )
        self.tabla_pagados.configure(
            xscrollcommand=scroll_h_pag.set, yscrollcommand=scroll_v_pag.set
        )

        scroll_v_pag.pack(side=tk.RIGHT, fill=tk.Y)
        self.tabla_pagados.pack(fill=tk.BOTH, expand=True)
        scroll_h_pag.pack(fill=tk.X)

        self._configurar_cabeceras_tablas(
            self.tabla_pagados, columnas_pag, es_pagado=True
        )

        frame_btns_pag = tk.Frame(self.tab_pagados, bg="#ffffff", pady=8, padx=5)
        frame_btns_pag.pack(fill=tk.X)

        tk.Button(
            frame_btns_pag,
            text="📂  Ver Voucher",
            bg="#6b46c1",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            command=self.ver_voucher_seleccionado,
            padx=6,
            pady=4,
        ).pack(side=tk.LEFT, padx=(0, 4))
        tk.Button(
            frame_btns_pag,
            text="📎  Adjuntar Voucher",
            bg="#2b6cb0",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            command=self.subir_voucher_tardio,
            padx=6,
            pady=4,
        ).pack(side=tk.LEFT, padx=4)
        tk.Button(
            frame_btns_pag,
            text="✏️  Editar Registro",
            bg="#dd6b20",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            command=self.abrir_editar_deuda_pagada,
            padx=6,
            pady=4,
        ).pack(side=tk.LEFT, padx=4)
        tk.Button(
            frame_btns_pag,
            text="❌  Eliminar",
            bg="#e53e3e",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            command=lambda: self.eliminar_obligacion(self.tabla_pagados),
            padx=6,
            pady=4,
        ).pack(side=tk.LEFT, padx=4)
        tk.Button(
            frame_btns_pag,
            text="📊  Exportar a Excel",
            bg="#234e52",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            command=self.exportar_pagados,
            padx=8,
            pady=4,
        ).pack(side=tk.RIGHT)

    def _configurar_cabeceras_tablas(self, tabla, columnas, es_pagado=False):
        for col in columnas:
            tabla.heading(col, text=col)

        tabla.column("ID", width=45, minwidth=40, anchor="center")
        tabla.column("Razon Social", width=190, minwidth=140)
        tabla.column("RUC", width=100, minwidth=90, anchor="center")
        tabla.column("Régimen", width=120, minwidth=100, anchor="center")
        tabla.column("Tipo", width=65, minwidth=55, anchor="center")
        tabla.column("Situación", width=100, minwidth=90, anchor="center")
        tabla.column("Periodo", width=75, minwidth=70, anchor="center")
        tabla.column("Monto", width=90, minwidth=80, anchor="e")
        tabla.column("Vencimiento", width=95, minwidth=85, anchor="center")

        if es_pagado:
            tabla.column("Fecha Pago", width=95, minwidth=85, anchor="center")
            tabla.column("Voucher", width=70, minwidth=60, anchor="center")
            tabla.column("Tributos", width=200, minwidth=150)
        else:
            tabla.column("Tributos", width=250, minwidth=180)

    # =====================================================================
    # 3. LÓGICA DE NEGOCIO Y GUARDADO DE ARCHIVOS
    # =====================================================================
    def autocalcular_vencimiento(self, event=None):
        sel_idx = self.cmb_clientes.current()
        periodo = self.ent_periodo.get().strip()
        if sel_idx != -1 and periodo and not self.ent_vence.get().strip():
            ruc = self.lista_clientes_db[sel_idx][1]
            venc_sugerido = calcular_vencimiento_sunat_aproximado(ruc, periodo)
            if venc_sugerido:
                self.ent_vence.delete(0, tk.END)
                self.ent_vence.insert(0, venc_sugerido)

    def _copiar_voucher_local(self, ruta_origen, id_obligacion):
        """Copia el archivo a la carpeta local ./vouchers/ para no depender de la ruta original."""
        if not ruta_origen or not os.path.exists(ruta_origen):
            return None
        _, ext = os.path.splitext(ruta_origen)
        nombre_destino = f"voucher_{id_obligacion}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
        ruta_destino = os.path.join(CARPETA_VOUCHERS, nombre_destino)
        shutil.copy(ruta_origen, ruta_destino)
        return ruta_destino

    def al_buscar(self, event):
        self.cargar_datos()

    def limpiar_filtros(self):
        self.cmb_filtro_ruc.current(0)
        self.ent_filtro_periodo.delete(0, tk.END)
        self.cargar_datos()

    def cargar_datos(self):
        for item in self.tabla_pendientes.get_children():
            self.tabla_pendientes.delete(item)
        for item in self.tabla_pagados.get_children():
            self.tabla_pagados.delete(item)

        texto_busqueda = self.ent_buscar.get().strip()
        idx_filtro_cliente = self.cmb_filtro_ruc.current()
        filtro_periodo = self.ent_filtro_periodo.get().strip()

        conn = conectar_bd()
        cursor = conn.cursor()

        query_pend = """
        SELECT o.id, c.razon_social, c.ruc, c.regimen, o.tipo, o.situacion, o.periodo, o.monto, o.fecha_vencimiento, o.descripcion_tributos
        FROM obligaciones o JOIN clientes c ON o.cliente_id = c.id WHERE o.estado = 'PENDIENTE'
        """
        params_pend = []
        if texto_busqueda:
            query_pend += " AND (c.razon_social LIKE ? OR c.ruc LIKE ?)"
            params_pend.extend([f"%{texto_busqueda}%", f"%{texto_busqueda}%"])
        if idx_filtro_cliente > 0:
            query_pend += " AND o.cliente_id = ?"
            params_pend.append(self.lista_clientes_db[idx_filtro_cliente - 1][0])
        if filtro_periodo:
            query_pend += " AND o.periodo LIKE ?"
            params_pend.append(f"%{filtro_periodo}%")
        query_pend += " ORDER BY o.fecha_vencimiento ASC"

        cursor.execute(query_pend, params_pend)
        hoy = datetime.now().date()
        limite_3dias = hoy + timedelta(days=3)

        for f in cursor.fetchall():
            venc_str = f[8]
            tag_alerta = ""
            if venc_str:
                try:
                    f_venc = datetime.strptime(venc_str, "%Y-%m-%d").date()
                    if f_venc < hoy:
                        tag_alerta = "vencido"
                    elif hoy <= f_venc <= limite_3dias:
                        tag_alerta = "por_vencer"
                except ValueError:
                    pass

            tags = (tag_alerta,) if tag_alerta else ()
            self.tabla_pendientes.insert(
                "",
                tk.END,
                values=(
                    f[0],
                    f[1],
                    f[2],
                    f[3],
                    f[4],
                    f[5],
                    f[6],
                    f"{f[7]:.2f}",
                    f[8] or "",
                    f[9] or "",
                ),
                tags=tags,
            )

        query_pag = """
        SELECT o.id, c.razon_social, c.ruc, c.regimen, o.tipo, o.situacion, o.periodo, o.monto, o.fecha_vencimiento, o.fecha_pago, o.ruta_voucher, o.descripcion_tributos
        FROM obligaciones o JOIN clientes c ON o.cliente_id = c.id WHERE o.estado = 'PAGADO'
        """
        params_pag = []
        if texto_busqueda:
            query_pag += " AND (c.razon_social LIKE ? OR c.ruc LIKE ?)"
            params_pag.extend([f"%{texto_busqueda}%", f"%{texto_busqueda}%"])
        if idx_filtro_cliente > 0:
            query_pag += " AND o.cliente_id = ?"
            params_pag.append(self.lista_clientes_db[idx_filtro_cliente - 1][0])
        if filtro_periodo:
            query_pag += " AND o.periodo LIKE ?"
            params_pag.append(f"%{filtro_periodo}%")
        query_pag += " ORDER BY o.fecha_pago DESC"

        cursor.execute(query_pag, params_pag)
        for f in cursor.fetchall():
            self.tabla_pagados.insert(
                "",
                tk.END,
                values=(
                    f[0],
                    f[1],
                    f[2],
                    f[3],
                    f[4],
                    f[5],
                    f[6],
                    f"{f[7]:.2f}",
                    f[8] or "",
                    f[9] or "",
                    "SÍ" if f[10] else "NO",
                    f[11] or "",
                ),
            )

        cursor.execute(
            "SELECT id, ruc, razon_social FROM clientes ORDER BY razon_social ASC"
        )
        self.lista_clientes_db = cursor.fetchall()

        act_sel = self.cmb_filtro_ruc.get()
        opciones = ["-- TODOS LOS CLIENTES --"] + [
            cl[2] for cl in self.lista_clientes_db
        ]
        self.cmb_filtro_ruc["values"] = opciones
        if act_sel in opciones:
            self.cmb_filtro_ruc.set(act_sel)
        else:
            self.cmb_filtro_ruc.current(0)

        self.cmb_clientes["values"] = [cl[2] for cl in self.lista_clientes_db]
        conn.close()

    def agregar_cliente(self):
        ruc = self.ent_ruc.get().strip()
        razon = self.ent_razon.get().strip()
        regimen = self.cmb_regimen.get()
        contacto = self.ent_contacto.get().strip()

        if not ruc or not razon:
            messagebox.showwarning(
                "Atención", "RUC y Razón Social son campos requeridos."
            )
            return
        if len(ruc) != 11 or not ruc.isdigit():
            messagebox.showwarning(
                "Atención", "El RUC debe tener exactamente 11 dígitos numéricos."
            )
            return

        try:
            conn = conectar_bd()
            cursor = conn.cursor()
            cursor.execute(
                "INSERT INTO clientes (ruc, razon_social, contacto, regimen) VALUES (?, ?, ?, ?)",
                (ruc, razon, contacto, regimen),
            )
            conn.commit()
            conn.close()

            messagebox.showinfo(
                "Éxito", f"Cliente '{razon}' registrado correctamente."
            )
            self.ent_ruc.delete(0, tk.END)
            self.ent_razon.delete(0, tk.END)
            self.ent_contacto.delete(0, tk.END)
            self.cargar_datos()
        except sqlite3.IntegrityError:
            messagebox.showerror(
                "Error", "Este número de RUC ya existe registrado."
            )

    def agregar_obligacion(self):
        sel_idx = self.cmb_clientes.current()
        tipo = self.cmb_tipo.get()
        situacion = self.cmb_situacion.get()
        periodo = self.ent_periodo.get().strip()
        monto_raw = self.ent_monto.get().strip()
        vence = self.ent_vence.get().strip()
        fecha_pago = self.ent_fecha_pago_reg.get().strip()
        tributos = self.ent_tributos.get().strip()

        if sel_idx == -1:
            messagebox.showwarning(
                "Atención", "Debe seleccionar un cliente de la lista."
            )
            return
        if not periodo or not monto_raw:
            messagebox.showwarning(
                "Atención", "El periodo y el monto son datos obligatorios."
            )
            return
        if not validar_periodo(periodo):
            messagebox.showwarning(
                "Atención",
                "El formato del periodo debe ser AAAA-MM (Ejemplo: 2026-07).",
            )
            return
        if vence and not validar_fecha_iso(vence):
            messagebox.showwarning(
                "Atención",
                "El formato de la fecha de vencimiento debe ser AAAA-MM-DD.",
            )
            return
        if fecha_pago and not validar_fecha_iso(fecha_pago):
            messagebox.showwarning(
                "Atención",
                "El formato de la fecha de pago debe ser AAAA-MM-DD.",
            )
            return

        try:
            monto = float(monto_raw)
            cliente_id = self.lista_clientes_db[sel_idx][0]
            estado = "PAGADO" if fecha_pago else "PENDIENTE"

            conn = conectar_bd()
            cursor = conn.cursor()
            cursor.execute(
                "INSERT INTO obligaciones (cliente_id, tipo, periodo, monto, fecha_vencimiento, estado, fecha_pago, situacion, descripcion_tributos) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    cliente_id,
                    tipo,
                    periodo,
                    monto,
                    vence,
                    estado,
                    fecha_pago if fecha_pago else None,
                    situacion,
                    tributos,
                ),
            )
            conn.commit()
            conn.close()

            messagebox.showinfo(
                "Éxito", "Obligación tributaria guardada correctamente."
            )
            self.ent_periodo.delete(0, tk.END)
            self.ent_monto.delete(0, tk.END)
            self.ent_vence.delete(0, tk.END)
            self.ent_fecha_pago_reg.delete(0, tk.END)
            self.ent_tributos.delete(0, tk.END)
            self.cargar_datos()
        except ValueError:
            messagebox.showerror(
                "Error", "Asegúrese de ingresar un monto numérico válido."
            )

    def marcar_como_pagado(self):
        seleccion = self.tabla_pendientes.selection()
        if not seleccion:
            messagebox.showwarning(
                "Atención", "Seleccione al menos una obligación de la lista."
            )
            return

        cant = len(seleccion)
        confirmar = messagebox.askyesno(
            "Confirmar Pago",
            f"¿Marcar como pagada(s) las {cant} obligación(es) seleccionada(s)?",
        )
        if not confirmar:
            return

        ruta_origen = None
        if cant == 1:
            if messagebox.askyesno(
                "Voucher",
                "¿Desea adjuntar la captura del voucher ahora mismo?",
            ):
                ruta_origen = filedialog.askopenfilename(
                    title="Seleccionar Voucher",
                    filetypes=[("Archivos de Imagen/PDF", "*.pdf *.png *.jpg *.jpeg")],
                )

        fecha_hoy = datetime.now().strftime("%Y-%m-%d")
        conn = conectar_bd()
        cursor = conn.cursor()

        for idx in seleccion:
            id_ob = self.tabla_pendientes.item(idx, "values")[0]
            ruta_guardada = None
            if cant == 1 and ruta_origen:
                ruta_guardada = self._copiar_voucher_local(ruta_origen, id_ob)

            cursor.execute(
                "UPDATE obligaciones SET estado='PAGADO', fecha_pago=?, ruta_voucher=? WHERE id=?",
                (
                    fecha_hoy,
                    ruta_guardada
                    if ruta_guardada
                    else (
                        None
                        if cant == 1
                        else cursor.execute(
                            "SELECT ruta_voucher FROM obligaciones WHERE id=?",
                            (id_ob,),
                        ).fetchone()[0]
                    ),
                    id_ob,
                ),
            )

        conn.commit()
        conn.close()
        self.cargar_datos()

    def ver_voucher_seleccionado(self):
        sel = self.tabla_pagados.selection()
        if not sel:
            messagebox.showwarning(
                "Atención", "Seleccione un registro del historial."
            )
            return
        id_ob = self.tabla_pagados.item(sel[0], "values")[0]

        conn = conectar_bd()
        cursor = conn.cursor()
        cursor.execute("SELECT ruta_voucher FROM obligaciones WHERE id=?", (id_ob,))
        res = cursor.fetchone()
        conn.close()

        if res and res[0] and os.path.exists(res[0]):
            if os.name == "nt":
                os.startfile(res[0])
            else:
                subprocess.run(
                    ["open" if os.name == "mac" else "xdg-open", res[0]]
                )
        else:
            messagebox.showerror(
                "No Encontrado",
                "No hay voucher asociado o el archivo no existe en el sistema.",
            )

    def subir_voucher_tardio(self):
        sel = self.tabla_pagados.selection()
        if not sel:
            return
        id_ob = self.tabla_pagados.item(sel[0], "values")[0]

        ruta_origen = filedialog.askopenfilename(
            title="Adjuntar Voucher",
            filetypes=[("Archivos de Imagen/PDF", "*.pdf *.png *.jpg *.jpeg")],
        )
        if ruta_origen:
            ruta_destino = self._copiar_voucher_local(ruta_origen, id_ob)
            conn = conectar_bd()
            cursor = conn.cursor()
            cursor.execute(
                "UPDATE obligaciones SET ruta_voucher=? WHERE id=?",
                (ruta_destino, id_ob),
            )
            conn.commit()
            conn.close()
            self.cargar_datos()

    def eliminar_obligacion(self, tabla):
        sel = tabla.selection()
        if not sel:
            return
        vals = tabla.item(sel[0], "values")
        if messagebox.askyesno(
            "Eliminar",
            f"¿Seguro que desea eliminar el registro ID {vals[0]} de {vals[1]}?",
        ):
            conn = conectar_bd()
            cursor = conn.cursor()
            cursor.execute("DELETE FROM obligaciones WHERE id=?", (vals[0],))
            conn.commit()
            conn.close()
            self.cargar_datos()

    # =====================================================================
    # MODALES DE EDICIÓN ESTILIZADOS
    # =====================================================================
    def abrir_editar_cliente(self):
        sel = self.cmb_clientes.current()
        if sel == -1:
            messagebox.showwarning(
                "Atención",
                "Seleccione un cliente en el menú desplegable izquierdo para editarlo.",
            )
            return
        id_c = self.lista_clientes_db[sel][0]

        conn = conectar_bd()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT ruc, razon_social, regimen, contacto FROM clientes WHERE id=?",
            (id_c,),
        )
        cl = cursor.fetchone()
        conn.close()

        win = tk.Toplevel(self.root)
        win.title("Editar Cliente")
        win.geometry("380x320")
        win.config(bg="#ffffff")
        win.grab_set()

        frame_win = tk.Frame(win, bg="#ffffff", padx=15, pady=15)
        frame_win.pack(fill=tk.BOTH, expand=True)

        tk.Label(
            frame_win, text="RUC:", bg="#ffffff", font=("Segoe UI", 9, "bold")
        ).pack(anchor="w")
        e_r = tk.Entry(frame_win, font=("Segoe UI", 10), bd=1, relief="solid")
        e_r.insert(0, cl[0])
        e_r.pack(fill=tk.X, pady=(2, 8))

        tk.Label(
            frame_win,
            text="Razón Social:",
            bg="#ffffff",
            font=("Segoe UI", 9, "bold"),
        ).pack(anchor="w")
        e_rz = tk.Entry(frame_win, font=("Segoe UI", 10), bd=1, relief="solid")
        e_rz.insert(0, cl[1])
        e_rz.pack(fill=tk.X, pady=(2, 8))

        tk.Label(
            frame_win, text="Régimen:", bg="#ffffff", font=("Segoe UI", 9, "bold")
        ).pack(anchor="w")
        c_rg = ttk.Combobox(
            frame_win,
            values=[
                "MYPE Tributario",
                "Régimen General",
                "RER (Especial)",
                "NRUS",
            ],
            state="readonly",
        )
        c_rg.set(cl[2])
        c_rg.pack(fill=tk.X, pady=(2, 8))

        tk.Label(
            frame_win, text="Contacto:", bg="#ffffff", font=("Segoe UI", 9, "bold")
        ).pack(anchor="w")
        e_ct = tk.Entry(frame_win, font=("Segoe UI", 10), bd=1, relief="solid")
        e_ct.insert(0, cl[3] or "")
        e_ct.pack(fill=tk.X, pady=(2, 15))

        def guardar():
            ruc_val = e_r.get().strip()
            if len(ruc_val) != 11 or not ruc_val.isdigit():
                messagebox.showerror(
                    "Error", "El RUC debe tener 11 dígitos numéricos."
                )
                return

            conn = conectar_bd()
            cursor = conn.cursor()
            try:
                cursor.execute(
                    "UPDATE clientes SET ruc=?, razon_social=?, regimen=?, contacto=? WHERE id=?",
                    (
                        ruc_val,
                        e_rz.get().strip(),
                        c_rg.get(),
                        e_ct.get().strip(),
                        id_c,
                    ),
                )
                conn.commit()
                win.destroy()
                self.cargar_datos()
            except sqlite3.IntegrityError:
                messagebox.showerror("Error", "RUC duplicado.")
            finally:
                conn.close()

        tk.Button(
            frame_win,
            text="Guardar Cambios",
            bg="#2b6cb0",
            fg="white",
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            command=guardar,
            pady=4,
        ).pack(fill=tk.X)

    def abrir_editar_deuda_pendiente(self):
        self._abrir_modal_obligacion(self.tabla_pendientes)

    def abrir_editar_deuda_pagada(self):
        self._abrir_modal_obligacion(self.tabla_pagados)

    def _abrir_modal_obligacion(self, tabla):
        sel = tabla.selection()
        if not sel:
            return
        id_ob = tabla.item(sel[0], "values")[0]

        conn = conectar_bd()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT tipo, situacion, periodo, monto, fecha_vencimiento, descripcion_tributos FROM obligaciones WHERE id=?",
            (id_ob,),
        )
        ob = cursor.fetchone()
        conn.close()

        win = tk.Toplevel(self.root)
        win.title("Modificar Obligación")
        win.geometry("380x420")
        win.config(bg="#ffffff")
        win.grab_set()

        frame_win = tk.Frame(win, bg="#ffffff", padx=15, pady=15)
        frame_win.pack(fill=tk.BOTH, expand=True)

        tk.Label(
            frame_win, text="Tipo:", bg="#ffffff", font=("Segoe UI", 9, "bold")
        ).pack(anchor="w")
        c_t = ttk.Combobox(frame_win, values=["SUNAT", "AFP"], state="readonly")
        c_t.set(ob[0])
        c_t.pack(fill=tk.X, pady=(2, 6))

        tk.Label(
            frame_win, text="Situación:", bg="#ffffff", font=("Segoe UI", 9, "bold")
        ).pack(anchor="w")
        c_s = ttk.Combobox(
            frame_win,
            values=["Declarado", "Por Declarar", "Fraccionado"],
            state="readonly",
        )
        c_s.set(ob[1])
        c_s.pack(fill=tk.X, pady=(2, 6))

        tk.Label(
            frame_win, text="Periodo:", bg="#ffffff", font=("Segoe UI", 9, "bold")
        ).pack(anchor="w")
        e_p = tk.Entry(frame_win, font=("Segoe UI", 10), bd=1, relief="solid")
        e_p.insert(0, ob[2])
        e_p.pack(fill=tk.X, pady=(2, 6))

        tk.Label(
            frame_win,
            text="Monto (S/):",
            bg="#ffffff",
            font=("Segoe UI", 9, "bold"),
        ).pack(anchor="w")
        e_m = tk.Entry(frame_win, font=("Segoe UI", 10), bd=1, relief="solid")
        e_m.insert(0, ob[3])
        e_m.pack(fill=tk.X, pady=(2, 6))

        tk.Label(
            frame_win,
            text="Vencimiento:",
            bg="#ffffff",
            font=("Segoe UI", 9, "bold"),
        ).pack(anchor="w")
        e_v = tk.Entry(frame_win, font=("Segoe UI", 10), bd=1, relief="solid")
        e_v.insert(0, ob[4] or "")
        e_v.pack(fill=tk.X, pady=(2, 6))

        tk.Label(
            frame_win,
            text="Detalle de Tributos:",
            bg="#ffffff",
            font=("Segoe UI", 9, "bold"),
        ).pack(anchor="w")
        e_tr = tk.Entry(frame_win, font=("Segoe UI", 10), bd=1, relief="solid")
        e_tr.insert(0, ob[5] or "")
        e_tr.pack(fill=tk.X, pady=(2, 15))

        def guardar():
            periodo = e_p.get().strip()
            vence = e_v.get().strip()
            if not validar_periodo(periodo):
                messagebox.showerror(
                    "Error", "El periodo debe tener el formato AAAA-MM."
                )
                return
            if vence and not validar_fecha_iso(vence):
                messagebox.showerror(
                    "Error", "La fecha de vencimiento debe tener formato AAAA-MM-DD."
                )
                return

            try:
                monto = float(e_m.get())
                conn = conectar_bd()
                cursor = conn.cursor()
                cursor.execute(
                    "UPDATE obligaciones SET tipo=?, situacion=?, periodo=?, monto=?, fecha_vencimiento=?, descripcion_tributos=? WHERE id=?",
                    (
                        c_t.get(),
                        c_s.get(),
                        periodo,
                        monto,
                        vence,
                        e_tr.get().strip(),
                        id_ob,
                    ),
                )
                conn.commit()
                conn.close()
                win.destroy()
                self.cargar_datos()
            except ValueError:
                messagebox.showerror(
                    "Error", "Asegúrese de ingresar un monto numérico válido."
                )

        tk.Button(
            frame_win,
            text="Actualizar Datos",
            bg="#2b6cb0",
            fg="white",
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            command=guardar,
            pady=4,
        ).pack(fill=tk.X)

    # =====================================================================
    # 4. EXPORTACIÓN A EXCEL FILTRADA DINÁMICAMENTE
    # =====================================================================
    def exportar_pendientes(self):
        if not OPENPYXL_DISPONIBLE:
            return
        self._guardar_excel_reporte(es_pago=False)

    def exportar_pagados(self):
        if not OPENPYXL_DISPONIBLE:
            return
        self._guardar_excel_reporte(es_pago=True)

    def _guardar_excel_reporte(self, es_pago=False):
        nombre_defecto = (
            "Reporte_Deudas_Pendientes"
            if not es_pago
            else "Reporte_Historial_Pagos"
        )
        nombre_hoja = "Pendientes" if not es_pago else "Pagados"

        ruta_archivo = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Libro de Excel", "*.xlsx")],
            initialfile=f"{nombre_defecto}_{datetime.now().strftime('%Y%m%d')}.xlsx",
            title="Guardar Reporte Excel",
        )
        if not ruta_archivo:
            return

        texto_busqueda = self.ent_buscar.get().strip()
        idx_filtro_cliente = self.cmb_filtro_ruc.current()
        filtro_periodo = self.ent_filtro_periodo.get().strip()

        conn = conectar_bd()
        cursor = conn.cursor()
        estado_filtro = "PAGADO" if es_pago else "PENDIENTE"

        query = f"""
        SELECT o.id, c.razon_social, c.ruc, c.regimen, o.tipo, o.situacion, o.periodo, o.monto, o.fecha_vencimiento
        {", o.fecha_pago" if es_pago else ""}, o.descripcion_tributos
        FROM obligaciones o JOIN clientes c ON o.cliente_id = c.id WHERE o.estado = '{estado_filtro}'
        """
        params = []
        if texto_busqueda:
            query += " AND (c.razon_social LIKE ? OR c.ruc LIKE ?)"
            params.extend([f"%{texto_busqueda}%", f"%{texto_busqueda}%"])
        if idx_filtro_cliente > 0:
            query += " AND o.cliente_id = ?"
            params.append(self.lista_clientes_db[idx_filtro_cliente - 1][0])
        if filtro_periodo:
            query += " AND o.periodo LIKE ?"
            params.append(f"%{filtro_periodo}%")

        query += (
            " ORDER BY o.fecha_vencimiento ASC"
            if not es_pago
            else " ORDER BY o.fecha_pago DESC"
        )

        cursor.execute(query, params)
        datos = cursor.fetchall()
        conn.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = nombre_hoja
        ws.views.sheetView[0].showGridLines = True

        color_navy = "1A365D"
        font_titulo = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        font_cabecera = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        font_datos = Font(name="Segoe UI", size=10)
        font_total = Font(name="Segoe UI", size=10, bold=True)

        fill_titulo = PatternFill(
            start_color=color_navy, end_color=color_navy, fill_type="solid"
        )
        fill_cabecera = PatternFill(
            start_color="2D3748", end_color="2D3748", fill_type="solid"
        )
        fill_zebra = PatternFill(
            start_color="F7FAFC", end_color="F7FAFC", fill_type="solid"
        )

        align_centro = Alignment(horizontal="center", vertical="center")
        align_derecha = Alignment(horizontal="right", vertical="center")
        align_izq = Alignment(horizontal="left", vertical="center")

        border_cell = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )
        border_total = Border(
            top=Side(style="thin", color="1A365D"),
            bottom=Side(style="double", color="1A365D"),
        )

        num_cols = 10 if not es_pago else 11
        max_letra = openpyxl.utils.get_column_letter(num_cols)

        ws.merge_cells(f"A1:{max_letra}1")
        ws["A1"] = (
            f"REPORTE DE OBLIGACIONES TRIBUTARIAS ({nombre_hoja.upper()})"
        )
        ws["A1"].font = font_titulo
        ws["A1"].fill = fill_titulo
        ws["A1"].alignment = align_centro
        ws.row_dimensions[1].height = 35

        ws.append([])

        headers = [
            "ID",
            "Razón Social",
            "RUC",
            "Régimen",
            "Tipo",
            "Situación",
            "Periodo",
            "Monto",
            "Vencimiento",
        ]
        if es_pago:
            headers.append("Fecha Pago")
        headers.append("Detalle Tributos")

        ws.append(headers)
        ws.row_dimensions[3].height = 24

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.font = font_cabecera
            cell.fill = fill_cabecera
            cell.alignment = align_centro
            cell.border = border_cell

        for idx, fila in enumerate(datos):
            ws.append(list(fila))
            r_idx = idx + 4
            ws.row_dimensions[r_idx].height = 20
            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = font_datos
                cell.border = border_cell
                if idx % 2 == 1:
                    cell.fill = fill_zebra

                if c_idx in [1, 3, 4, 5, 6, 7, 9] or (c_idx == 10 and es_pago):
                    cell.alignment = align_centro
                elif c_idx == 8:
                    cell.alignment = align_derecha
                    cell.number_format = '"S/"#,##0.00'
                else:
                    cell.alignment = align_izq

        tot_row = len(datos) + 4
        ws.cell(row=tot_row, column=7, value="TOTAL:").font = font_total
        ws.cell(row=tot_row, column=7).alignment = align_derecha

        c_sum = ws.cell(row=tot_row, column=8, value=f"=SUM(H4:H{tot_row-1})")
        c_sum.font = font_total
        c_sum.alignment = align_derecha
        c_sum.number_format = '"S/"#,##0.00'
        c_sum.border = border_total

        for col in ws.columns:
            max_len = max(
                len(str(cell.value or "")) for cell in col if cell.row > 2
            )
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

        wb.save(ruta_archivo)
        messagebox.showinfo(
            "Éxito",
            "El archivo Excel se generó correctamente con los filtros seleccionados.",
        )


if __name__ == "__main__":
    root = tk.Tk()
    app = AppControlDeudas(root)
    root.mainloop()