import os
import re
from datetime import datetime, timedelta
import io
import streamlit as st
from supabase import create_client, Client

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    OPENPYXL_DISPONIBLE = True
except ImportError:
    OPENPYXL_DISPONIBLE = False

# =====================================================================
# CONFIGURACIÓN Y ESTILOS INSPIRADOS EN SUNAT
# =====================================================================
st.set_page_config(
    page_title="Sistema de Gestión - Estilo SUNAT",
    layout="wide"
)

st.markdown("""
    <style>
    [data-testid="InputInstructions"] { display: none !important; }
    .main { background-color: #F4F6F9; font-family: 'Segoe UI', Arial, sans-serif; }
    .titulo-corporativo {
        color: #003366; font-size: 24px; font-weight: 700; letter-spacing: -0.3px;
        margin-bottom: 20px; padding-bottom: 10px; border-bottom: 3px solid #D91A2A;
    }
    .subtitulo-corporativo {
        color: #003366; font-size: 15px; font-weight: 700; margin-top: 15px;
        margin-bottom: 10px; text-transform: uppercase; letter-spacing: 0.5px;
    }
    [data-testid="stSidebar"] { background-color: #FFFFFF; border-right: 1px solid #E0E0E0; }
    [data-testid="stSidebar"] p, [data-testid="stSidebar"] label, [data-testid="stSidebar"] span { color: #333333 !important; }
    [data-testid="stSidebar"] input { color: #000000 !important; background-color: #FFFFFF !important; border: 1px solid #CCCCCC !important; }
    div[data-baseweb="select"] > div { background-color: #FFFFFF !important; color: #000000 !important; border: 1px solid #CCCCCC !important; }
    div[data-baseweb="select"] span { color: #000000 !important; }
    div[data-baseweb="icon"] svg { fill: #003366 !important; }
    .stTabs [data-baseweb="tab-list"] { gap: 6px; background-color: #E9ECEF; padding: 6px; border-radius: 4px; }
    .stTabs [data-baseweb="tab"] { height: 38px; white-space: pre-wrap; border-radius: 4px; font-weight: 600; color: #003366; }
    .stTabs [aria-selected="true"] { background-color: #D91A2A !important; color: #FFFFFF !important; }
    div[data-testid="stFormSubmitButton"] > button, div.stButton > button {
        background-color: #003366 !important; color: #FFFFFF !important;
        font-weight: 700 !important; font-size: 15px !important; border-radius: 4px !important;
        border: none !important; width: 100% !important;
    }
    div[data-testid="stFormSubmitButton"] > button p, div.stButton > button p { color: #FFFFFF !important; }
    div[data-testid="stFormSubmitButton"] > button:hover, div.stButton > button:hover { background-color: #D91A2A !important; }
    </style>
""", unsafe_allow_html=True)

# =====================================================================
# 1. CONEXIÓN A SUPABASE
# =====================================================================
@st.cache_resource
def get_supabase_client() -> Client:
    url = st.secrets["SUPABASE_URL"]
    key = st.secrets["SUPABASE_KEY"]
    return create_client(url, key)

supabase = get_supabase_client()

CARPETA_VOUCHERS = "vouchers"
if not os.path.exists(CARPETA_VOUCHERS):
    os.makedirs(CARPETA_VOUCHERS)

def validar_periodo(periodo_str):
    if not periodo_str:
        return False
    return bool(re.match(r"^\d{4}-(0[1-9]|1[0-2])$", periodo_str))

def calcular_vencimiento_sunat_aproximado(ruc, periodo):
    if not ruc or len(ruc) != 11 or not validar_periodo(periodo):
        return ""
    try:
        ultimo_digito = int(ruc[-1])
        año, mes = map(int, periodo.split("-"))
        if mes == 12:
            mes_venc = 1
            año_venc = año + 1
        else:
            mes_venc = mes + 1
            año_venc = año

        dias_offset = {0: 14, 1: 15, 2: 16, 3: 17, 4: 18, 5: 19, 6: 20, 7: 21, 8: 22, 9: 23}
        dia = dias_offset.get(ultimo_digito, 15)
        return f"{año_venc:04d}-{mes_venc:02d}-{dia:02d}"
    except Exception:
        return ""

def guardar_voucher_local(uploaded_file, id_obligacion):
    if uploaded_file is None:
        return None
    ext = os.path.splitext(uploaded_file.name)[1]
    nombre_destino = f"voucher_{id_obligacion}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
    ruta_destino = os.path.join(CARPETA_VOUCHERS, nombre_destino)
    with open(ruta_destino, "wb") as f:
        f.write(uploaded_file.getbuffer())
    return ruta_destino

def generar_excel_bytes(es_pago, texto_busqueda, id_cliente_filtro, filtro_periodo):
    estado_filtro = "PAGADO" if es_pago else "PENDIENTE"
    query = supabase.table("obligaciones").select("id, tipo, situacion, periodo, monto, fecha_vencimiento, fecha_pago, descripcion_tributos, clientes(id, razon_social, ruc, regimen)").eq("estado", estado_filtro)

    if id_cliente_filtro and id_cliente_filtro != 0:
        query = query.eq("cliente_id", id_cliente_filtro)
    if filtro_periodo:
        query = query.ilike("periodo", f"%{filtro_periodo}%")

    res = query.order("fecha_vencimiento" if not es_pago else "fecha_pago", desc=es_pago).execute()
    filas = res.data or []

    datos = []
    for item in filas:
        cli = item.get("clientes") or {}
        razon = cli.get("razon_social", "")
        ruc = cli.get("ruc", "")
        regimen = cli.get("regimen", "")

        if texto_busqueda:
            t = texto_busqueda.lower()
            if t not in razon.lower() and t not in ruc.lower():
                continue

        fila = [item["id"], razon, ruc, regimen, item["tipo"], item["situacion"], item["periodo"], float(item["monto"]), item["fecha_vencimiento"]]
        if es_pago:
            fila.append(item.get("fecha_pago", ""))
        fila.append(item.get("descripcion_tributos", ""))
        datos.append(fila)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pagados" if es_pago else "Pendientes"
    ws.views.sheetView[0].showGridLines = True

    color_navy = "003366"
    fill_titulo = PatternFill(start_color=color_navy, end_color=color_navy, fill_type="solid")
    fill_cabecera = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    font_titulo = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    font_cabecera = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_datos = Font(name="Segoe UI", size=10)
    font_total = Font(name="Segoe UI", size=10, bold=True)

    align_centro = Alignment(horizontal="center", vertical="center")
    align_derecha = Alignment(horizontal="right", vertical="center")
    align_izq = Alignment(horizontal="left", vertical="center")

    border_cell = Border(left=Side(style="thin", color="E2E8F0"), right=Side(style="thin", color="E2E8F0"),
                         top=Side(style="thin", color="E2E8F0"), bottom=Side(style="thin", color="E2E8F0"))
    border_total = Border(top=Side(style="thin", color="003366"), bottom=Side(style="double", color="003366"))

    headers = ["ID", "Razón Social", "RUC", "Régimen", "Tipo", "Situación", "Periodo", "Monto", "Vencimiento"]
    if es_pago:
        headers.append("Fecha Pago")
    headers.append("Detalle Tributos")

    max_letra = openpyxl.utils.get_column_letter(len(headers))
    ws.merge_cells(f"A1:{max_letra}1")
    ws["A1"] = f"REPORTE DE OBLIGACIONES TRIBUTARIAS - STATUS: {ws.title.upper()}"
    ws["A1"].font = font_titulo
    ws["A1"].fill = fill_titulo
    ws["A1"].alignment = align_centro
    ws.row_dimensions[1].height = 35

    ws.append([])
    ws.append(headers)
    ws.row_dimensions[3].height = 24

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = font_cabecera
        cell.fill = fill_cabecera
        cell.alignment = align_centro
        cell.border = border_cell

    for idx, fila in enumerate(datos):
        ws.append(list(fila))
        r_idx = idx + 4
        ws.row_dimensions[r_idx].height = 20
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = font_datos
            cell.border = border_cell
            if idx % 2 == 1:
                cell.fill = fill_zebra

            if c_idx in [1, 3, 4, 5, 6, 7, 9] or (c_idx == 10 and es_pago):
                cell.alignment = align_centro
            elif c_idx == 8:
                cell.alignment = align_derecha
                cell.number_format = '"S/"#,##0.00'
            else:
                cell.alignment = align_izq

    tot_row = len(datos) + 4
    ws.cell(row=tot_row, column=7, value="TOTAL GENERAL:").font = font_total
    ws.cell(row=tot_row, column=7).alignment = align_derecha

    c_sum = ws.cell(row=tot_row, column=8, value=f"=SUM(H4:H{tot_row-1})")
    c_sum.font = font_total
    c_sum.alignment = align_derecha
    c_sum.number_format = '"S/"#,##0.00'
    c_sum.border = border_total

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col if cell.row > 2)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

# =====================================================================
# 2. INTERFAZ PRINCIPAL
# =====================================================================
st.markdown('<div class="titulo-corporativo">SUNAT - SISTEMA DE CONTROL DE OBLIGACIONES Y COMPROBANTES</div>', unsafe_allow_html=True)

# Cargar catálogo de clientes desde Supabase
res_cli = supabase.table("clientes").select("id, ruc, razon_social").order("razon_social").execute()
clientes_db = [(c["id"], c["ruc"], c["razon_social"]) for c in res_cli.data] if res_cli.data else []

dic_clientes = {cl[0]: cl[2] for cl in clientes_db}
dic_ruc = {cl[0]: cl[1] for cl in clientes_db}

# BARRA LATERAL
st.sidebar.markdown('<div style="font-size: 15px; font-weight: 700; color: #003366; margin-bottom: 15px; border-bottom: 2px solid #D91A2A; padding-bottom: 5px;">MÓDULO DE REGISTRO</div>', unsafe_allow_html=True)

opcion_sidebar = st.sidebar.radio("Seleccione Operación:", ["Registro de Contribuyente", "Registro de Obligación"])

if opcion_sidebar == "Registro de Contribuyente":
    st.sidebar.markdown('<div class="subtitulo-corporativo">Datos del Contribuyente</div>', unsafe_allow_html=True)
    
    with st.sidebar.form(key="form_contribuyente", clear_on_submit=True):
        reg_ruc = st.text_input("Número de RUC (11 dígitos):", max_chars=11)
        reg_razon = st.text_input("Razón Social / Nombre Comercial:")
        reg_regimen = st.selectbox("Régimen Tributario:", ["MYPE Tributario", "Régimen General", "RER (Especial)", "NRUS"])
        reg_contacto = st.text_input("Contacto Administrativo:")
        btn_guardar_cli = st.form_submit_button("Guardar Contribuyente", use_container_width=True)

    if btn_guardar_cli:
        if len(reg_ruc) != 11 or not reg_ruc.isdigit():
            st.sidebar.error("El RUC debe contar con exactamente 11 dígitos numéricos.")
        elif not reg_razon.strip():
            st.sidebar.error("El campo Razón Social es obligatorio.")
        else:
            try:
                supabase.table("clientes").insert({
                    "ruc": reg_ruc,
                    "razon_social": reg_razon,
                    "contacto": reg_contacto,
                    "regimen": reg_regimen
                }).execute()
                st.sidebar.success(f"Contribuyente '{reg_razon}' registrado correctamente.")
                st.rerun()
            except Exception as e:
                st.sidebar.error("Error al registrar: verifique que el RUC no esté duplicado.")

elif opcion_sidebar == "Registro de Obligación":
    st.sidebar.markdown('<div class="subtitulo-corporativo">Detalle de la Obligación</div>', unsafe_allow_html=True)
    if not clientes_db:
        st.sidebar.warning("No existen contribuyentes registrados. Ingrese uno antes de continuar.")
    else:
        cliente_sel = st.sidebar.selectbox("Contribuyente:", options=list(dic_clientes.keys()), format_func=lambda x: dic_clientes[x])
        tipo_ob = st.sidebar.selectbox("Entidad / Tipo:", ["SUNAT", "AFP"])
        situacion_ob = st.sidebar.selectbox("Estado de Declaración:", ["Declarado", "Por Declarar", "Fraccionado"])
        periodo_ob = st.sidebar.text_input("Periodo Tributario (AAAA-MM):", placeholder="Ej: 2026-07")
        
        ruc_actual = dic_ruc.get(cliente_sel, "")
        venc_auto = calcular_vencimiento_sunat_aproximado(ruc_actual, periodo_ob) if tipo_ob == "SUNAT" else ""

        with st.sidebar.form(key="form_obligacion", clear_on_submit=True):
            venc_ob = st.text_input("Fecha de Vencimiento (AAAA-MM-DD):", value=venc_auto, placeholder="AAAA-MM-DD")
            monto_ob = st.number_input("Monto Determinado (S/):", min_value=0.0, step=10.0, format="%.2f")
            pago_ob = st.date_input("Fecha de Pago Ejecutado (Opcional):", value=None, format="DD/MM/YYYY")
            tributos_ob = st.text_input("Descripción / Código de Tributo:")
            btn_guardar_ob = st.form_submit_button("Guardar Obligación", use_container_width=True)

        if btn_guardar_ob:
            if not validar_periodo(periodo_ob):
                st.sidebar.error("El periodo debe mantener la estructura AAAA-MM.")
            elif monto_ob <= 0:
                st.sidebar.error("El monto ingresado debe ser mayor a 0.00.")
            else:
                estado_ob = "PAGADO" if pago_ob else "PENDIENTE"
                fecha_pago_str = pago_ob.strftime("%Y-%m-%d") if pago_ob else None
                supabase.table("obligaciones").insert({
                    "cliente_id": cliente_sel,
                    "tipo": tipo_ob,
                    "periodo": periodo_ob,
                    "monto": monto_ob,
                    "fecha_vencimiento": venc_ob,
                    "estado": estado_ob,
                    "fecha_pago": fecha_pago_str,
                    "situacion": situacion_ob,
                    "descripcion_tributos": tributos_ob
                }).execute()
                st.sidebar.success("Obligación financiera registrada en el sistema.")
                st.rerun()

# PANEL DE FILTROS
st.markdown('<div class="subtitulo-corporativo">Filtros de Búsqueda y Consulta</div>', unsafe_allow_html=True)
col_filtro1, col_filtro2, col_filtro3 = st.columns([2, 2, 1])

with col_filtro1:
    txt_buscar = st.text_input("Búsqueda por Razón Social o RUC:", placeholder="Ingrese texto de búsqueda...")
with col_filtro2:
    opciones_filtro_cliente = {0: "TODOS LOS CONTRIBUYENTES"}
    opciones_filtro_cliente.update(dic_clientes)
    cliente_filtro_id = st.selectbox("Filtrar por Contribuyente:", options=list(opciones_filtro_cliente.keys()), format_func=lambda x: opciones_filtro_cliente[x])
with col_filtro3:
    periodo_filtro = st.text_input("Filtrar por Periodo:", placeholder="AAAA-MM")

tab_pendientes, tab_pagados = st.tabs(["Obligaciones Pendientes", "Historial de Obligaciones Pagadas"])

# ----------------- TABLA DE PENDIENTES -----------------
with tab_pendientes:
    q_pend = supabase.table("obligaciones").select("id, tipo, situacion, periodo, monto, fecha_vencimiento, descripcion_tributos, clientes(id, razon_social, ruc, regimen)").eq("estado", "PENDIENTE")
    if cliente_filtro_id != 0:
        q_pend = q_pend.eq("cliente_id", cliente_filtro_id)
    if periodo_filtro:
        q_pend = q_pend.ilike("periodo", f"%{periodo_filtro}%")

    res_pend = q_pend.order("fecha_vencimiento", desc=False).execute()
    pend_data_raw = res_pend.data or []

    filas_pend = []
    for item in pend_data_raw:
        cli = item.get("clientes") or {}
        razon = cli.get("razon_social", "")
        ruc = cli.get("ruc", "")
        regimen = cli.get("regimen", "")

        if txt_buscar:
            t = txt_buscar.lower()
            if t not in razon.lower() and t not in ruc.lower():
                continue

        filas_pend.append((
            item["id"], razon, ruc, regimen, item["tipo"], item["situacion"],
            item["periodo"], float(item["monto"]), item.get("fecha_vencimiento") or "",
            item.get("descripcion_tributos") or ""
        ))

    if filas_pend:
        data_pend = []
        hoy = datetime.now().date()
        limite = hoy + timedelta(days=3)

        for f in filas_pend:
            alerta = "En Plazo"
            if f[8]:
                try:
                    f_venc = datetime.strptime(f[8], "%Y-%m-%d").date()
                    if f_venc < hoy:
                        alerta = "Vencido"
                    elif hoy <= f_venc <= limite:
                        alerta = "Próximo a Vencer"
                except ValueError:
                    pass

            data_pend.append({
                "ID": f[0], "Estatus": alerta, "Razón Social": f[1], "RUC": f[2], "Régimen": f[3],
                "Tipo": f[4], "Situación": f[5], "Periodo": f[6], "Monto (S/)": f"{f[7]:.2f}",
                "Vencimiento": f[8], "Detalle / Tributos": f[9]
            })

        st.dataframe(data_pend, use_container_width=True)

        col_accion1, col_accion2 = st.columns([1, 1])
        with col_accion1:
            st.markdown('<div class="subtitulo-corporativo">Procesar Pago de Obligación</div>', unsafe_allow_html=True)
            id_pagar = st.number_input("ID de Obligación a Liquidar:", min_value=1, step=1)
            voucher_file = st.file_uploader("Adjuntar Comprobante de Pago (Imagen/PDF):", type=["png", "jpg", "jpeg", "pdf"])

            if st.button("Confirmar Pago", use_container_width=True):
                check = supabase.table("obligaciones").select("id").eq("id", id_pagar).eq("estado", "PENDIENTE").execute()
                if check.data:
                    ruta_v = guardar_voucher_local(voucher_file, id_pagar)
                    hoy_str = datetime.now().strftime("%Y-%m-%d")
                    supabase.table("obligaciones").update({"estado": "PAGADO", "fecha_pago": hoy_str, "ruta_voucher": ruta_v}).eq("id", id_pagar).execute()
                    st.success(f"La obligación con ID {id_pagar} fue actualizada a estado PAGADO.")
                    st.rerun()
                else:
                    st.error("El ID especificado no corresponde a una obligación pendiente válida.")

        with col_accion2:
            st.markdown('<div class="subtitulo-corporativo">Eliminación de Registro</div>', unsafe_allow_html=True)
            id_elim = st.number_input("ID de Obligación a Remover:", min_value=1, step=1)
            if st.button("Eliminar Registro", use_container_width=True):
                supabase.table("obligaciones").delete().eq("id", id_elim).execute()
                st.warning(f"La obligación con ID {id_elim} ha sido removida del sistema.")
                st.rerun()

        # MODIFICAR OBLIGACIÓN PENDIENTE
        st.markdown('<div class="subtitulo-corporativo">Modificar Obligación Pendiente</div>', unsafe_allow_html=True)
        
        lista_opciones_editar = [0] + [f[0] for f in filas_pend]
        dic_opciones_label = {
            0: "-- Seleccione una obligación para editar --",
            **{f[0]: f"ID {f[0]} | {f[1]} | Periodo: {f[6]} | S/ {f[7]:.2f}" for f in filas_pend}
        }

        id_mod = st.selectbox(
            "Seleccione la Obligación a Modificar:",
            options=lista_opciones_editar,
            format_func=lambda x: dic_opciones_label.get(x, str(x)),
            key="mod_id_select"
        )

        if id_mod != 0:
            res_ob = supabase.table("obligaciones").select("*").eq("id", id_mod).eq("estado", "PENDIENTE").execute()
            if res_ob.data:
                ob_data = res_ob.data[0]
                with st.form(key=f"form_modificar_ob_{id_mod}"):
                    col_m1, col_m2 = st.columns(2)
                    with col_m1:
                        lista_tipos = ["SUNAT", "AFP"]
                        idx_tipo = lista_tipos.index(ob_data["tipo"]) if ob_data["tipo"] in lista_tipos else 0
                        m_tipo = st.selectbox("Entidad / Tipo:", lista_tipos, index=idx_tipo)

                        lista_sit = ["Declarado", "Por Declarar", "Fraccionado"]
                        idx_sit = lista_sit.index(ob_data["situacion"]) if ob_data["situacion"] in lista_sit else 0
                        m_situacion = st.selectbox("Estado de Declaración:", lista_sit, index=idx_sit)

                        m_periodo = st.text_input("Periodo Tributario (AAAA-MM):", value=ob_data["periodo"])

                    with col_m2:
                        m_monto = st.number_input("Monto Determinado (S/):", min_value=0.01, step=10.0, format="%.2f", value=float(ob_data["monto"]))
                        m_venc = st.text_input("Fecha de Vencimiento (AAAA-MM-DD):", value=ob_data.get("fecha_vencimiento") or "")
                        m_tributos = st.text_input("Descripción / Código de Tributo:", value=ob_data.get("descripcion_tributos") or "")

                    btn_guardar_mod = st.form_submit_button("Guardar Cambios", use_container_width=True)

                if btn_guardar_mod:
                    if not validar_periodo(m_periodo):
                        st.error("El periodo debe mantener la estructura AAAA-MM.")
                    elif m_monto <= 0:
                        st.error("El monto ingresado debe ser mayor a 0.00.")
                    else:
                        supabase.table("obligaciones").update({
                            "tipo": m_tipo,
                            "situacion": m_situacion,
                            "periodo": m_periodo,
                            "monto": m_monto,
                            "fecha_vencimiento": m_venc,
                            "descripcion_tributos": m_tributos
                        }).eq("id", id_mod).execute()
                        st.success(f"La obligación ID {id_mod} se actualizó exitosamente.")
                        st.rerun()

        if OPENPYXL_DISPONIBLE:
            st.markdown("---")
            bytes_excel = generar_excel_bytes(False, txt_buscar, cliente_filtro_id, periodo_filtro)
            st.download_button("Exportar Reporte de Pendientes (Excel)", data=bytes_excel, file_name=f"Reporte_Pendientes_{datetime.now().strftime('%Y%m%d')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.info("No se registran obligaciones pendientes con los criterios seleccionados.")

# ----------------- TABLA DE PAGADOS -----------------
with tab_pagados:
    q_pag = supabase.table("obligaciones").select("id, tipo, situacion, periodo, monto, fecha_vencimiento, fecha_pago, ruta_voucher, descripcion_tributos, clientes(id, razon_social, ruc, regimen)").eq("estado", "PAGADO")
    if cliente_filtro_id != 0:
        q_pag = q_pag.eq("cliente_id", cliente_filtro_id)
    if periodo_filtro:
        q_pag = q_pag.ilike("periodo", f"%{periodo_filtro}%")

    res_pag = q_pag.order("fecha_pago", desc=True).execute()
    pag_data_raw = res_pag.data or []

    filas_pag = []
    for item in pag_data_raw:
        cli = item.get("clientes") or {}
        razon = cli.get("razon_social", "")
        ruc = cli.get("ruc", "")
        regimen = cli.get("regimen", "")

        if txt_buscar:
            t = txt_buscar.lower()
            if t not in razon.lower() and t not in ruc.lower():
                continue

        filas_pag.append((
            item["id"], razon, ruc, regimen, item["tipo"], item["situacion"],
            item["periodo"], float(item["monto"]), item.get("fecha_vencimiento") or "",
            item.get("fecha_pago") or "", item.get("ruta_voucher"), item.get("descripcion_tributos") or ""
        ))

    if filas_pag:
        data_pag = []
        for f in filas_pag:
            data_pag.append({
                "ID": f[0], "Razón Social": f[1], "RUC": f[2], "Régimen": f[3],
                "Tipo": f[4], "Situación": f[5], "Periodo": f[6], "Monto (S/)": f"{f[7]:.2f}",
                "Vencimiento": f[8], "Fecha Pago": f[9], "Comprobante": "Adjunto" if f[10] else "Sin Archivo",
                "Detalle / Tributos": f[11]
            })

        st.dataframe(data_pag, use_container_width=True)

        col_v1, col_v2 = st.columns(2)
        with col_v1:
            st.markdown('<div class="subtitulo-corporativo">Visualización de Comprobante</div>', unsafe_allow_html=True)
            id_ver_v = st.number_input("ID de Registro a Consultar:", min_value=1, step=1, key="v_ver")
            if st.button("Consultar Comprobante", use_container_width=True):
                r_v = supabase.table("obligaciones").select("ruta_voucher").eq("id", id_ver_v).execute()
                if r_v.data and r_v.data[0].get("ruta_voucher") and os.path.exists(r_v.data[0]["ruta_voucher"]):
                    ruta_archivo = r_v.data[0]["ruta_voucher"]
                    extension = os.path.splitext(ruta_archivo)[1].lower()
                    
                    if extension in [".png", ".jpg", ".jpeg"]:
                        st.image(ruta_archivo, use_container_width=True)
                    elif extension == ".pdf":
                        with open(ruta_archivo, "rb") as f:
                            pdf_bytes = f.read()
                        st.success("Comprobante en formato PDF disponible.")
                        st.download_button(
                            label="Descargar Comprobante PDF",
                            data=pdf_bytes,
                            file_name=os.path.basename(ruta_archivo),
                            mime="application/pdf",
                            use_container_width=True
                        )
                else:
                    st.error("No se registra comprobante adjunto para el ID indicado.")

        with col_v2:
            st.markdown('<div class="subtitulo-corporativo">Actualización de Comprobante</div>', unsafe_allow_html=True)
            id_sub_v = st.number_input("ID de Registro Objetivo:", min_value=1, step=1, key="v_sub")
            v_tardio = st.file_uploader("Seleccionar Nuevo Archivo:", type=["png", "jpg", "jpeg", "pdf"], key="f_sub")
            if st.button("Actualizar Archivo", use_container_width=True):
                if v_tardio:
                    ruta_t = guardar_voucher_local(v_tardio, id_sub_v)
                    supabase.table("obligaciones").update({"ruta_voucher": ruta_t}).eq("id", id_sub_v).execute()
                    st.success("Comprobante adjuntado correctamente.")
                    st.rerun()

        if OPENPYXL_DISPONIBLE:
            st.markdown("---")
            bytes_excel_p = generar_excel_bytes(True, txt_buscar, cliente_filtro_id, periodo_filtro)
            st.download_button("Exportar Historial de Pagados (Excel)", data=bytes_excel_p, file_name=f"Reporte_Pagados_{datetime.now().strftime('%Y%m%d')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.info("No se registran transacciones finalizadas en el historial.")